"""Export extraction results to a formatted Excel (.xlsx) workbook."""
from __future__ import annotations

import os
from datetime import datetime
from typing import List, Optional

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from .parser import InvoiceFields

HEADERS = [
    "#",
    "File",
    "Status",
    "Invoice No",
    "Invoice Date",
    "Vendor / Supplier",
    "GSTIN",
    "Taxable Amount",
    "CGST",
    "SGST",
    "IGST",
    "Total Invoice Amount",
    "Method / Engine",
    "Notes",
]
WIDTHS = [4, 34, 13, 18, 14, 36, 20, 15, 12, 12, 12, 18, 22, 40]
MONEY_COLS = {8, 9, 10, 11, 12}
MONEY_FORMAT = "#,##0.00"

_HEADER_FILL = PatternFill("solid", fgColor="1F4E79")
_HEADER_FONT = Font(bold=True, color="FFFFFF", size=11)
_ERROR_FILL = PatternFill("solid", fgColor="FCE4E4")
_STRIPE_FILL = PatternFill("solid", fgColor="F2F7FB")
_BORDER = Border(
    left=Side(style="thin", color="BFBFBF"),
    right=Side(style="thin", color="BFBFBF"),
    top=Side(style="thin", color="BFBFBF"),
    bottom=Side(style="thin", color="BFBFBF"),
)

__all__ = ["HEADERS", "build_row", "export_results"]


def build_row(
    path: str,
    *,
    fields: Optional[InvoiceFields] = None,
    status: str = "OK",
    method: str = "",
    engine: str = "",
    note: str = "",
) -> dict:
    """Build one output row (dict keyed by :data:`HEADERS`)."""
    f = fields
    return {
        "File": os.path.basename(path),
        "Status": status,
        "Invoice No": f.invoice_no if f else "",
        "Invoice Date": f.date if f else "",
        "Vendor / Supplier": f.vendor if f else "",
        "GSTIN": f.gstin if f else "",
        "Taxable Amount": f.taxable if f else None,
        "CGST": f.cgst if f else None,
        "SGST": f.sgst if f else None,
        "IGST": f.igst if f else None,
        "Total Invoice Amount": f.total if f else None,
        "Method / Engine": " / ".join(x for x in (method, engine) if x),
        "Notes": note,
    }


def export_results(rows: List[dict], path: str) -> str:
    """Write *rows* to *path* as a formatted workbook. Returns the path."""
    wb = Workbook()
    ws = wb.active
    ws.title = "Invoices"

    # ---- header row -------------------------------------------------------
    ws.append(HEADERS)
    for col in range(1, len(HEADERS) + 1):
        cell = ws.cell(row=1, column=col)
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = _BORDER
        ws.column_dimensions[get_column_letter(col)].width = WIDTHS[col - 1]
    ws.row_dimensions[1].height = 30

    # ---- data rows ----------------------------------------------------------
    for idx, row in enumerate(rows, start=1):
        ws.append([idx] + [row.get(h, "") for h in HEADERS[1:]])
        r = ws.max_row
        for col in range(1, len(HEADERS) + 1):
            cell = ws.cell(row=r, column=col)
            cell.border = _BORDER
            if col in MONEY_COLS and isinstance(cell.value, (int, float)):
                cell.number_format = MONEY_FORMAT
                cell.alignment = Alignment(horizontal="right")
            elif col == 1:
                cell.alignment = Alignment(horizontal="center")
        if str(row.get("Status", "")).lower() in {"error", "no data found"}:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=col).fill = _ERROR_FILL
        elif idx % 2 == 0:
            for col in range(1, len(HEADERS) + 1):
                ws.cell(row=r, column=col).fill = _STRIPE_FILL

    ws.freeze_panes = "A2"
    if len(rows) >= 1:
        ws.auto_filter.ref = f"A1:{get_column_letter(len(HEADERS))}{ws.max_row}"

    # ---- summary sheet --------------------------------------------------------
    s = wb.create_sheet("Summary")
    s.column_dimensions["A"].width = 30
    s.column_dimensions["B"].width = 60
    ok_rows = [r for r in rows if str(r.get("Status", "")).lower() not in {"error", "no data found"}]
    s["A1"] = "Invoice Extraction Summary"
    s["A1"].font = Font(bold=True, size=14, color="1F4E79")
    s["A3"] = "Generated on"
    s["B3"] = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    s["A4"] = "Total files processed"
    s["B4"] = len(rows)
    s["A5"] = "Files with data"
    s["B5"] = len(ok_rows)
    s["A6"] = "Errors / no data"
    s["B6"] = len(rows) - len(ok_rows)
    methods = sorted({r.get("Method / Engine", "") for r in rows if r.get("Method / Engine")})
    s["A7"] = "Methods / engines used"
    s["B7"] = ", ".join(m for m in methods if m) or "-"
    s["A9"] = "Note"
    s["B9"] = ("Always verify extracted values against the original invoices before "
               "entering them into your accounts.")
    for row in s.iter_rows(min_row=3, max_row=9, min_col=1, max_col=1):
        for cell in row:
            cell.font = Font(bold=True)
    for row in s.iter_rows(min_row=9, max_row=9, min_col=2, max_col=2):
        for cell in row:
            cell.alignment = Alignment(wrap_text=True)

    os.makedirs(os.path.dirname(os.path.abspath(path)) or ".", exist_ok=True)
    wb.save(path)
    return path
