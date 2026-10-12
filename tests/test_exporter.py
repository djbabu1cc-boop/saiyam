"""Tests for the Excel exporter."""
import os
import tempfile
import unittest

from openpyxl import load_workbook

from invoice_extractor.exporter import HEADERS, build_row, export_results
from invoice_extractor.parser import InvoiceFields


class TestExporter(unittest.TestCase):
    def _rows(self):
        fields = InvoiceFields(
            invoice_no="INV/2026/0042",
            date="2026-10-05",
            vendor="ACME SUPPLIERS PVT LTD",
            gstin="06AAACA1234A1Z5",
            taxable=26000.0,
            cgst=2340.0,
            sgst=2340.0,
            igst=None,
            total=30680.0,
        )
        rows = [
            build_row("invoices/invoice_1_gst.pdf", fields=fields,
                      status="OK", method="pdf-text", engine=""),
            build_row("invoices/broken.pdf", status="Error",
                      note="FileNotFoundError: file not found"),
        ]
        return rows

    def test_export_and_read_back(self):
        rows = self._rows()
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.xlsx")
            export_results(rows, path)
            self.assertTrue(os.path.exists(path))

            wb = load_workbook(path)
            self.assertEqual(wb.sheetnames[0], "Invoices")
            self.assertIn("Summary", wb.sheetnames)

            ws = wb["Invoices"]
            headers = [c.value for c in ws[1]]
            self.assertEqual(headers, HEADERS)
            self.assertEqual(ws.max_row, 3)  # header + 2 rows

            # first data row
            self.assertEqual(ws.cell(row=2, column=headers.index("Invoice No") + 1).value,
                             "INV/2026/0042")
            taxable_cell = ws.cell(row=2, column=headers.index("Taxable Amount") + 1)
            self.assertEqual(taxable_cell.value, 26000.0)
            self.assertEqual(taxable_cell.number_format, "#,##0.00")

            # error row
            status_col = headers.index("Status") + 1
            self.assertEqual(ws.cell(row=3, column=status_col).value, "Error")

    def test_summary_counts(self):
        with tempfile.TemporaryDirectory() as tmp:
            path = os.path.join(tmp, "out.xlsx")
            export_results(self._rows(), path)
            wb = load_workbook(path)
            s = wb["Summary"]
            values = {s.cell(row=r, column=1).value: s.cell(row=r, column=2).value
                      for r in range(1, s.max_row + 1)}
            self.assertEqual(values.get("Total files processed"), 2)
            self.assertEqual(values.get("Files with data"), 1)
            self.assertEqual(values.get("Errors / no data"), 1)


if __name__ == "__main__":
    unittest.main()
